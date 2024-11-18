__coding__ = "utf-8"

import logging

from docx import Document

from pojo.MSTReqPOJO import ReqPOJO
from tools.common.dat_csv_common import create_file_path
from tools.report.report_generation import replace_placeholders_in_boa_tables
from tools.utils.DateUtils import get_current_date

logging.basicConfig(level=logging.DEBUG, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')

def replace_mst_header_page_docx(config_data: ReqPOJO, request_data: dict) -> str:
    replacements:dict = {}
    replacements["{1}"] = request_data["system"]
    replacements["{2}"] = request_data["vehicle_identity_number"]
    replacements["{3}"] = request_data["ecu_variant"]
    replacements["{4}"] = request_data["type_of_vehicle"]
    replacements["{5}"] = request_data["fuel_type"]
    replacements["{6}"] = request_data["mileage"]
    replacements["{7}"] = request_data["project_version"]
    replacements["{8}"] = request_data["gearbox"]
    replacements["{9}"] = request_data["accelerator"]
    replacements["{10}"] = request_data["emission_standard"]
    replacements["{11}"] = request_data["ecu_identification"]
    replacements["{12}"] = request_data["a2l_file"]
    replacements["{13}"] = request_data["hex_file"]
    replacements["{14}"] = request_data["tested_by"]
    replacements["{15}"] = request_data["dept"]
    replacements["{16}"] = get_current_date()

    template_name = config_data.template_name
    template_path = config_data.template_path

    template_file_name, template_path = create_file_path(template_name, 'docx', template_path, 'template')
    logging.info("MST_TEST Header Page:%s,%s", template_path, template_file_name)

    output_path = config_data.output_path
    doc_output_name = config_data.doc_output_name
    doc_name, output_path = create_file_path(doc_output_name, 'docx', output_path, 'docx')

    doc = Document(template_path)
    replace_placeholders_in_boa_tables(doc, replacements)
    doc.save(output_path)

    return output_path
